cat > /home/claude/telegram-menu-bot/handlers/backup.py << 'PYEOF'
import io

from aiogram import F, Router
from aiogram.types import BufferedInputFile, CallbackQuery
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import database as db
import keyboards as kb
from utils import render

router = Router()

HEADER_FILL = PatternFill(start_color="FFD54F", end_color="FFD54F", fill_type="solid")
HEADER_FONT = Font(bold=True, size=12)
TITLE_FONT = Font(bold=True, size=14)


def _build_workbook(rows) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    grouped: dict[str, list] = {}
    for r in rows:
        grouped.setdefault(r["category"], []).append(r)

    if not grouped:
        ws = wb.create_sheet("Taomlar")
        ws.append(["Hozircha taomlar yo'q"])
    else:
        for category, items in grouped.items():
            sheet_name = (category or "Bo'lim")[:31]
            ws = wb.create_sheet(sheet_name)

            ws.merge_cells("A1:B1")
            ws["A1"] = f"🍽 {category}"
            ws["A1"].font = TITLE_FONT

            ws.append(["Taom nomi", "Retsept"])
            for cell in ws[2]:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL

            for item in items:
                ws.append([item["name"], item["recipe"] or ""])

            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 65
            for row in ws.iter_rows(min_row=3):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@router.callback_query(F.data == "backup")
async def backup_menu(call: CallbackQuery):
    cats = await db.get_categories()
    await render(call, "📦 Qaysi bo'limni Excel qilib olamiz?", kb.backup_scope_kb(cats))
    await call.answer()


@router.callback_query(F.data.startswith("backup_export:"))
async def backup_export(call: CallbackQuery):
    scope = call.data.split(":")[1]
    category_id = None if scope == "all" else int(scope)
    rows = await db.get_export_data(category_id)
    if not rows:
        await call.answer("Bu bo'limda hali taom yo'q.", show_alert=True)
        return

    data = _build_workbook(rows)
    filename = "taomlar_backup.xlsx" if category_id is None else f"{rows[0]['category']}.xlsx"
    file = BufferedInputFile(data, filename=filename)
    await call.message.answer_document(file, caption="📦 Excel fayl tayyor.")
    await call.answer()
PYEOF
python3 -c "import ast; ast.parse(open('/home/claude/telegram-menu-bot/handlers/backup.py').read()); print('OK')"
cat /home/claude/telegram-menu-bot/handlers/backup.py
